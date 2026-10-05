import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.contrib import admin, messages
from django.http import HttpResponse
from django.core.exceptions import ValidationError
from .models import contacto
from .forms import ContactoForm

# Personalización del encabezado y títulos del Admin
admin.site.site_header = "💼 Panel de Control Corporativo"
admin.site.site_title = "Agenda VIP Admin"
admin.site.index_title = "Panel de Gestión y Contactos"
admin.site.index_template = "admin/agenda_index.html"


@admin.action(description="📥 Exportar contactos seleccionados a Excel (.xlsx)")
def exportar_contactos_a_excel(modeladmin, request, queryset):
    """
    Genera un archivo .xlsx real con cada campo en su propia casilla,
    con encabezados estilizados y columnas auto-ajustadas.
    """
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Contactos"

        # Estilos para encabezados
        header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin', color='D4AF37'),
            right=Side(style='thin', color='D4AF37'),
            top=Side(style='thin', color='D4AF37'),
            bottom=Side(style='thin', color='D4AF37'),
        )

        # Encabezados
        headers = ['ID', 'Nombre Completo', 'Teléfono', 'Correo Electrónico', 'Dirección']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Datos
        data_font = Font(name='Calibri', size=10)
        data_align = Alignment(vertical='center')
        count = 0
        for row_num, obj in enumerate(queryset, 2):
            values = [obj.id, obj.nombre, obj.telefono, obj.correo, obj.direccion]
            for col_num, value in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border
            count += 1

        # Auto-ajustar anchos de columna
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_length + 4

        # Escribir a bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="contactos_admin.xlsx"'

        messages.success(request, f'✅ ¡Exportación exitosa! {count} contacto(s) exportados a Excel.')
        return response
    except Exception as ex:
        messages.error(request, f'❌ Error al generar el archivo Excel: {str(ex)}')
        return None


@admin.register(contacto)
class ContactoAdmin(admin.ModelAdmin):
    form = ContactoForm
    
    # Columnas mostradas en el listado
    list_display = ('id', 'nombre', 'telefono', 'correo', 'direccion')
    list_display_links = ('id', 'nombre')
    
    # Edición rápida desde la tabla
    list_editable = ('telefono', 'correo', 'direccion')
    
    # Barra de búsqueda universal
    search_fields = ('nombre', 'correo', 'telefono', 'direccion')
    
    # Filtros laterales
    list_filter = ('nombre',)
    
    # Paginación para cumplir con el requisito de interfaz
    list_per_page = 10
    list_max_show_all = 100
    
    # Orden predeterminado
    ordering = ('-id',)
    
    # Acciones personalizadas
    actions = [exportar_contactos_a_excel]

    # Inyección de CSS personalizado centrado y moderno
    class Media:
        css = {
            'all': ('admin/css/custom.css',)
        }

    def save_model(self, request, obj, form, change):
        """
        Manejo de errores y validación al guardar registros en el Admin.
        """
        try:
            super().save_model(request, obj, form, change)
            accion = "actualizado" if change else "creado"
            messages.success(request, f'✨ Contacto "{obj.nombre}" {accion} con éxito en el sistema.')
        except ValidationError as ve:
            messages.error(request, f'⚠️ Error de validación: {ve}')
            raise ve
        except Exception as e:
            messages.error(request, f'❌ Error inesperado al guardar contacto: {str(e)}')
            raise e

    def delete_model(self, request, obj):
        """
        Feedback al eliminar un contacto individual desde el Admin.
        """
        nombre = obj.nombre
        try:
            super().delete_model(request, obj)
            messages.warning(request, f'🗑️ El contacto "{nombre}" fue eliminado del panel administrativo.')
        except Exception as e:
            messages.error(request, f'❌ Error al eliminar el contacto: {str(e)}')
            raise e
